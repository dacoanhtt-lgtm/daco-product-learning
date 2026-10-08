# -*- coding: utf-8 -*-
"""
Script Đồng Bộ Dữ Liệu Tự Động 1-Click
Từ file Excel 'Nghiên cứu sản phẩm - Tuấn Anh.xlsx' sang thư mục 'data/' của Web Portal
"""
import sys
import io
import os
import json
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

excel_path = os.path.abspath(r'e:\sp\Mitutoyo\Nghiên cứu sản phẩm - Tuấn Anh.xlsx')
data_dir = os.path.abspath(r'e:\sp\Mitutoyo\cam_nang_san_pham\data')

print(f"[*] Dang mo file Excel: {excel_path}")
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Đồng bộ Mitsubishi
if 'SP Mitsubishi (Chuẩn 2026)' in wb.sheetnames:
    print("[*] Dang doc sheet 'SP Mitsubishi (Chuẩn 2026)'...")
    ws = wb['SP Mitsubishi (Chuẩn 2026)']
    
    # Đọc existing mitsubishi.json để giữ lại software và history
    mitsu_json_path = os.path.join(data_dir, 'mitsubishi.json')
    with open(mitsu_json_path, 'r', encoding='utf-8') as f:
        mitsu_data = json.load(f)
    
    products = []
    # Dữ liệu từ row 5 đến 29
    for r in range(5, 30):
        stt = ws.cell(r, 1).value
        cat = ws.cell(r, 2).value
        subcat = ws.cell(r, 3).value
        serial = ws.cell(r, 4).value
        status = ws.cell(r, 5).value or ''
        fake = ws.cell(r, 6).value or 'Không'
        renew = ws.cell(r, 7).value or 'Không'
        models_str = ws.cell(r, 8).value or ''
        alt_str = ws.cell(r, 9).value or ''
        specs_str = ws.cell(r, 10).value or ''
        notes_str = ws.cell(r, 11).value or ''
        vendor_str = ws.cell(r, 12).value or ''

        if not serial:
            continue

        models_list = [m.strip() for m in str(models_str).split('\n') if m.strip()]
        
        # Chuyển specs và notes thành bullet points
        points = []
        if specs_str:
            for line in str(specs_str).split('\n'):
                if line.strip():
                    points.append(f"**Thông số & Ứng dụng**: {line.strip()}")
        if notes_str:
            for line in str(notes_str).split('\n'):
                if line.strip():
                    points.append(f"**Lưu ý kỹ thuật**: {line.strip()}")

        # Tìm existing product để giữ software/brochure nếu có
        existing_p = next((p for p in mitsu_data.get('products', []) if p['id'] == stt), None)
        software = existing_p.get('software') if existing_p else ''
        brochure = existing_p.get('brochure') if existing_p else ''

        products.append({
            "id": stt,
            "cat": cat,
            "subcat": subcat,
            "serial": serial,
            "status": status,
            "fake": fake,
            "renew": renew,
            "models": models_list,
            "replacement": alt_str,
            "software": software,
            "brochure": brochure,
            "points": points,
            "suppliers": vendor_str
        })

    mitsu_data['products'] = products
    with open(mitsu_json_path, 'w', encoding='utf-8') as f:
        json.dump(mitsu_data, f, ensure_ascii=False, indent=2)
    print(f"    -> Da cap nhat {len(products)} san pham vao mitsubishi.json!")

print("[V] Hoan tat dong bo du lieu thanh cong!")
