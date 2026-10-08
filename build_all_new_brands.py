import openpyxl, sys, json, re, os

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('e:/sp/Mitutoyo/Theo dõi hỏi hàng & tiến độ hàng_backup.xlsx', data_only=True)

def clean(v):
    if v is None: return ''
    return str(v).strip()

def is_stop_row(val):
    if not val: return False
    vl = val.lower()
    return any(k in vl for k in ['1. tồn kho', 'tiêu chí', 'đánh giá', 'nội dung', '1. sale', '2. mua hàng'])

# --- 1. OMRON ---
def parse_omron():
    ws = wb['SP OMRON']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'PLC'
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        if is_stop_row(c0): break
        if not c0 and not clean(r[1]) and not clean(r[5]): continue
        if c0 and ('TỔNG' in c0.upper() or 'NỘI DUNG' in c0.upper()): continue
        
        if c0 and len(c0) > 1 and not clean(r[1]) and not clean(r[5]):
            current_cat = c0
            continue
        elif c0:
            current_cat = c0

        model_name = clean(r[1])
        if not model_name: continue
        
        status_raw = clean(r[2])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'hiếm' in status_raw.lower(): status = 'Hiếm'
        elif 'mới' in status_raw.lower() or 'new' in status_raw.lower(): status = 'Mới'

        fake = 'Có' if 'yes' in clean(r[3]).lower() else 'Không'
        renew = 'Có' if 'yes' in clean(r[4]).lower() else 'Không'
        
        models_raw = clean(r[5])
        models = [m.strip() for m in re.split(r'[,;\n]+', models_raw) if m.strip()]
        if not models and model_name:
            models = [model_name]

        suppliers = clean(r[6])
        web = clean(r[7])

        points = []
        if status_raw and status_raw.lower() != status.lower():
            points.append(f"**Đặc điểm**: {status_raw}")
        if models:
            points.append(f"**Mã thông dụng**: {', '.join(models[:5])}")
        if fake == 'Có':
            points.append("⚠️ Cảnh báo thực chiến: Thị trường có nguy cơ hàng FAKE / Bo mạch copy nhái.")
        if renew == 'Có':
            points.append("🔄 Cảnh báo thực chiến: Có rủi ro hàng cũ tháo máy Renew dựng lại vỏ.")
        points.append(f"**Thế mạnh NCC**: {suppliers or 'Hợp Long, Bảo An stock sẵn'}")

        products.append({
            "id": pid,
            "cat": current_cat,
            "subcat": f"{current_cat} Omron chính hãng",
            "serial": f"Omron {model_name}",
            "status": status,
            "fake": fake,
            "renew": renew,
            "models": models,
            "replacement": "Chuyển sang dòng CP2E thay thế CP1E" if status == 'Ngừng sx' and 'CP' in model_name else "",
            "points": points,
            "software": "CX-One, Sysmac Studio" if any(k in current_cat.lower() for k in ['plc', 'biến tần', 'servo']) else "Không dùng",
            "brochure": web if web.startswith('http') else "https://www.ia.omron.com/",
            "suppliers": suppliers or "HỢP LONG, BẢO AN"
        })
        pid += 1
    return {
        "brand": "Omron",
        "products": products,
        "history": [
            {
                "category": "Hệ Thống PLC & Bộ Điều Khiển Omron",
                "steps": [
                    {
                        "era": "1990 - 2005",
                        "name": "C200H / CQM1",
                        "status": "Đã khai tử",
                        "badge": "discontinued",
                        "software": "SYSWIN / CX-P cũ",
                        "highlight": "Dòng PLC lập trình kinh điển đời đầu trong các dây chuyền lắp ráp xe máy, dệt may."
                    },
                    {
                        "era": "2006 - 2022",
                        "name": "CP1E / CP1L / CP1H / CJ2M",
                        "status": "Thông dụng",
                        "badge": "classic",
                        "software": "CX-One (CX-Programmer)",
                        "highlight": "Dòng PLC phổ thông phổ biến bậc nhất nhà máy Việt Nam, tích hợp cổng USB/Ethernet."
                    },
                    {
                        "era": "2020 - Nay",
                        "name": "CP2E / NX1P2 / NJ Series",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "Sysmac Studio",
                        "highlight": "Kiến trúc IoT đồng bộ hóa Motion, Logic, An toàn, kết nối cơ sở dữ liệu SQL/OPC UA."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "CX-One",
                "version": "v4.60+",
                "target": "PLC CP1E, CP1L, CP1H, CJ1, CJ2, CS1, HMI NB, Servo G5",
                "purpose": "Bộ phần mềm tổng hợp tích hợp toàn diện: CX-Programmer, CX-Integrator, CX-Designer...",
                "note": "Bộ công cụ tiêu chuẩn hàng đầu phục vụ bảo trì, lập trình máy móc Omron hiện hữu.",
                "link": "https://www.ia.omron.com/products/family/1629/"
            },
            {
                "name": "Sysmac Studio",
                "version": "v1.54+",
                "target": "PLC Machine Controller NX1P2, NX102, NJ301, NJ501, HMI NA",
                "purpose": "Môi trường phát triển tích hợp (IDE) duy nhất cho Logic, Motion, Vision, Robot, 3D Simulation.",
                "note": "Nền tảng tự động hóa cao cấp nhất của Omron, hỗ trợ giao thức mạng EtherCAT tốc độ cao.",
                "link": "https://www.ia.omron.com/products/family/3116/"
            },
            {
                "name": "NB-Designer",
                "version": "v1.50+",
                "target": "Màn hình HMI cảm ứng Omron dòng NB: NB3Q, NB5Q, NB7W, NB10W",
                "purpose": "Thiết kế giao diện đồ họa HMI cảm ứng, gán tag địa chỉ PLC Omron/Mitsubishi/Siemens.",
                "note": "Phần mềm miễn phí, nhẹ nhàng, dễ sử dụng cho các ứng dụng hiển thị điều khiển cơ bản.",
                "link": "https://www.ia.omron.com/products/family/3120/"
            }
        ],
        "brochures": [
            {
                "title": "Omron Industrial Automation General Products Guide",
                "category": "Catalog Tổng Hợp",
                "desc": "Tổng hợp dải sản phẩm điều khiển tự động hóa, cảm biến, rơ le đóng cắt Omron toàn cầu.",
                "link": "https://www.ia.omron.com/products/",
                "isLocal": False
            },
            {
                "title": "Omron CP2E Series Micro PLC Brochure",
                "category": "Catalog Sản Phẩm",
                "desc": "Dòng PLC nhỏ gọn thông minh kết nối mạng kép 2 cổng Ethernet thế hệ mới thay thế CP1E.",
                "link": "https://www.ia.omron.com/products/family/3748/",
                "isLocal": False
            },
            {
                "title": "Omron E2B Proximity Sensors Catalog",
                "category": "Catalog Sản Phẩm",
                "desc": "Cảm biến tiệm cận kim loại thân trụ đồng thau chuẩn IP67 kinh tế và thông dụng nhất.",
                "link": "https://www.ia.omron.com/products/family/3133/",
                "isLocal": False
            }
        ]
    }

# --- 2. AUTONICS ---
def parse_autonics():
    ws = wb['SP Autonics']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'Cảm biến tiệm cận'
    pid = 1
    for r in rows[2:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        c2 = clean(r[2])
        if is_stop_row(c0) or is_stop_row(c1): break
        
        if c0 and not c2 and not clean(r[5]):
            current_cat = c0
            if c1: current_cat = f"{c0} - {c1}"
            continue
        elif c0:
            current_cat = c0

        model_name = c2 or c1
        if not model_name or 'TÊN' in model_name.upper(): continue

        status_raw = clean(r[3])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'new' in status_raw.lower() or 'mới' in status_raw.lower(): status = 'Mới'

        fake = 'Có' if 'yes' in clean(r[4]).lower() else 'Không'
        
        models_raw = clean(r[5])
        models = [m.strip() for m in re.split(r'[,;\n]+', models_raw) if m.strip()]
        if not models and model_name: models = [model_name]

        desc = clean(r[6])
        specs = clean(r[7]) if len(r) > 7 else ''
        suppliers = clean(r[8]) if len(r) > 8 else ''

        points = []
        if desc: points.append(f"**Đặc điểm nổi bật**: {desc}")
        if specs: points.append(f"**Thông số/Quy cách**: {specs}")
        if fake == 'Có': points.append("⚠️ Cảnh báo thực chiến: Thị trường có hàng FAKE / nhái thương hiệu.")
        points.append(f"**Nhà cung cấp**: {suppliers or 'Hợp Long, Bảo An, DACO'}")

        products.append({
            "id": pid,
            "cat": current_cat.split(' - ')[0] if ' - ' in current_cat else current_cat,
            "subcat": desc or current_cat,
            "serial": f"Autonics {model_name}",
            "status": status,
            "fake": fake,
            "renew": 'Không',
            "models": models,
            "replacement": "",
            "points": points,
            "software": "DAQMaster" if any(k in current_cat.lower() for k in ['nhiệt', 'bộ đếm', 'timer', 'áp suất']) else "Không dùng",
            "brochure": "https://www.autonics.com/product/category/all",
            "suppliers": suppliers or "Hợp Long, Bảo An, DACO"
        })
        pid += 1
    return {
        "brand": "Autonics",
        "products": products,
        "history": [
            {
                "category": "Cảm Biến & Bộ Điều Khiển Autonics",
                "steps": [
                    {
                        "era": "1995 - 2010",
                        "name": "Dòng PR / TZN / FX",
                        "status": "Kinh điển",
                        "badge": "classic",
                        "software": "Không dùng phần mềm",
                        "highlight": "Đặt nền móng thiết bị cảm biến và đồng hồ nhiệt độ giá cạnh tranh xuất xứ Hàn Quốc tại VN."
                    },
                    {
                        "era": "2010 - 2020",
                        "name": "Dòng PRFA / TK / CT",
                        "status": "Thông dụng",
                        "badge": "current",
                        "software": "DAQMaster v2",
                        "highlight": "Ra mắt cảm biến tiệm cận chống tia lửa hàn PRFA, bộ điều khiển nhiệt TK lấy mẫu 50ms."
                    },
                    {
                        "era": "2020 - Nay",
                        "name": "Dòng PRFD-K / TX / Modbus",
                        "status": "Chuẩn hiện hành",
                        "badge": "future",
                        "software": "DAQMaster v3.5+",
                        "highlight": "Tích hợp truyền thông công nghiệp, cảm biến vỏ kim loại toàn phần PRFD-K siêu bền."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "DAQMaster",
                "version": "v3.5+",
                "target": "Bộ điều khiển nhiệt độ TK, TC, TX, Bộ đếm Timer CT, Bộ hiển thị xung",
                "purpose": "Phần mềm thu thập dữ liệu (SCADA mini), vẽ biểu đồ nhiệt độ thời gian thực, sao lưu tham số.",
                "note": "Kết nối máy tính qua bộ chuyển đổi USB-RS485 SCM-US485 của Autonics.",
                "link": "https://www.autonics.com/support/download/software"
            }
        ],
        "brochures": [
            {
                "title": "Autonics Sensors & Controllers Overview Catalog",
                "category": "Catalog Tổng Hợp",
                "desc": "Tổng hợp cảm biến quang, cảm biến tiệm cận, cảm biến áp suất, bộ điều khiển nhiệt độ Autonics.",
                "link": "https://www.autonics.com/product/category/all",
                "isLocal": False
            }
        ]
    }

# --- 3. PATLITE ---
def parse_patlite():
    ws = wb['SP Patlite']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'Đèn Tháp tín hiệu'
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        if is_stop_row(c0) or is_stop_row(c1): break
        if not c1: 
            if c0: current_cat = c0
            continue
        if c0: current_cat = c0

        model_name = c1
        status_raw = clean(r[2])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'không thông dụng' in status_raw.lower() or 'hiếm' in status_raw.lower(): status = 'Hiếm'
        elif 'mới' in status_raw.lower(): status = 'Mới'

        feature = clean(r[3])
        models_raw = clean(r[4])
        models = [m.strip() for m in re.split(r'[,;\n]+', models_raw) if m.strip()]
        if not models: models = [model_name]

        specs = clean(r[5])
        app = clean(r[6])
        suppliers = clean(r[7])
        web = clean(r[8])

        points = []
        if feature: points.append(f"**Chức năng**: {feature}")
        if specs: points.append(f"**Thông số**: {specs}")
        if app: points.append(f"**Ứng dụng**: {app}")
        points.append(f"**Nhà cung cấp**: {suppliers or 'DACO, Kanetsu'}")

        products.append({
            "id": pid,
            "cat": current_cat,
            "subcat": feature or specs or current_cat,
            "serial": f"Patlite {model_name}",
            "status": status,
            "fake": 'Không',
            "renew": 'Không',
            "models": models,
            "replacement": "Chuyển sang dòng LR4/LR5/LR6 thay thế LME/LCE đã ngừng sản xuất" if status == 'Ngừng sx' else "",
            "points": points,
            "software": "Patlite Sound Editor" if 'loa' in current_cat.lower() or 'mp3' in feature.lower() else "Không dùng",
            "brochure": web if web.startswith('http') else "https://www.patlite.com/",
            "suppliers": suppliers or "DACO, Kanetsu, Kaleido"
        })
        pid += 1
    return {
        "brand": "Patlite",
        "products": products,
        "history": [
            {
                "category": "Đèn Tháp & Còi Báo Động Patlite",
                "steps": [
                    {
                        "era": "1995 - 2016",
                        "name": "LCE / LME / SKH",
                        "status": "Đã khai tử",
                        "badge": "discontinued",
                        "software": "Không dùng phần mềm",
                        "highlight": "Dòng đèn tháp kinh điển vỏ đục LME và đèn quay cơ học SKH gắn liền dây chuyền Nhật Bản."
                    },
                    {
                        "era": "2017 - Nay",
                        "name": "Dòng LR Series (LR4/LR5/LR6/LR7)",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "Không dùng phần mềm",
                        "highlight": "Thấu kính lăng trụ phân tán ánh sáng 360 độ siêu sáng, mô-đun xoay lắp ghép chuẩn IP65."
                    },
                    {
                        "era": "2021 - Nay",
                        "name": "Dòng Báo Mạng NHV / NE-IL",
                        "status": "Thế hệ mới",
                        "badge": "future",
                        "software": "Web GUI / SNMP / MQTT",
                        "highlight": "Báo hiệu qua mạng LAN Ethernet, hỗ trợ gửi email cảnh báo, thông báo thoại MP3 và IoT."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "Patlite Sound Editor",
                "version": "v2.1+",
                "target": "Còi báo MP3 BSV, BDV, BKV và đèn tháp thông minh NHV Series",
                "purpose": "Nạp file âm thanh MP3, cài đặt giọng nói hướng dẫn an toàn, cấu hình IP mạng cho đèn.",
                "note": "Sử dụng thẻ nhớ SD hoặc cáp USB cắm trực tiếp vào thiết bị Patlite.",
                "link": "https://www.patlite.com/download/"
            }
        ],
        "brochures": [
            {
                "title": "Patlite Visual & Audible Signaling Devices Global Catalog",
                "category": "Catalog Tổng Hợp",
                "desc": "Catalog tổng thể đèn tháp, đèn quay, đèn nhấp nháy, còi báo động công nghiệp Patlite Nhật Bản.",
                "link": "https://www.patlite.com/",
                "isLocal": False
            }
        ]
    }

# --- 4. BROTHER ---
def parse_brother():
    ws = wb['SP BROTHER']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'MÁY IN NHÃN CẦM TAY'
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        if is_stop_row(c0) or is_stop_row(c1): break
        if not c1:
            if c0: current_cat = c0
            continue
        if c0: current_cat = c0

        model_name = c1
        status_raw = clean(r[2])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'không thông dụng' in status_raw.lower(): status = 'Hiếm'

        fake = 'Có' if 'yes' in clean(r[3]).lower() else 'Không'
        renew = 'Có' if 'yes' in clean(r[4]).lower() else 'Không'
        desc = clean(r[5])
        suppliers = clean(r[6])
        web = clean(r[7])

        points = []
        if desc:
            for line in desc.split('\n'):
                if line.strip(): points.append(line.strip())
        if not points:
            points.append(f"Máy in nhãn chuyên dụng {model_name}")
        points.append(f"**Nhà cung cấp**: {suppliers or 'DACO, VP TECH, METAKING'}")

        products.append({
            "id": pid,
            "cat": current_cat,
            "subcat": "Máy in nhãn P-Touch & Ống lồng chính hãng Brother",
            "serial": f"Brother {model_name}",
            "status": status,
            "fake": fake,
            "renew": renew,
            "models": [model_name],
            "replacement": "",
            "points": points,
            "software": "P-touch Editor",
            "brochure": web if web.startswith('http') else "https://www.brother.com.vn/",
            "suppliers": suppliers or "DACO, VP TECH, METAKING, KHUÊ TÚ"
        })
        pid += 1
    return {
        "brand": "Brother",
        "products": products,
        "history": [
            {
                "category": "Máy In Nhãn & Ống Lồng Đầu Cốt Brother",
                "steps": [
                    {
                        "era": "2000 - 2015",
                        "name": "PT-E100 / PT-7600",
                        "status": "Đã khai tử",
                        "badge": "discontinued",
                        "software": "Bàn phím tích hợp",
                        "highlight": "Máy in nhãn cầm tay dã ngoại cho thợ điện và kỹ sư viễn thông."
                    },
                    {
                        "era": "2016 - Nay",
                        "name": "PT-E110 / PT-E850TKW",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "P-touch Editor v5",
                        "highlight": "Máy in ống lồng đầu cốt PT-E850TKW tích hợp 2 động cơ in nhãn và in ống đồng thời."
                    },
                    {
                        "era": "2023 - Nay",
                        "name": "PT-E310BT / PT-E560BT",
                        "status": "Thế hệ mới",
                        "badge": "future",
                        "software": "Brother Pro Label Tool App / Bluetooth",
                        "highlight": "In nhãn điều khiển trực tiếp từ smartphone qua Bluetooth, dao cắt Half-Cut tự động."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "P-touch Editor",
                "version": "v5.4+",
                "target": "Toàn bộ máy in nhãn Brother PT-E850TKW, PT-P900W, QL-800, TD-4420",
                "purpose": "Thiết kế nhãn mã vạch 1D/2D, bảng nhảy số tự động, liên kết dữ liệu trực tiếp với file Excel.",
                "note": "Phần mềm miễn phí chính hãng Brother cho máy tính Windows / macOS.",
                "link": "https://support.brother.com/g/b/downloadtop.aspx"
            }
        ],
        "brochures": [
            {
                "title": "Brother P-Touch Industrial Labeling Solutions Catalog",
                "category": "Catalog Tổng Hợp",
                "desc": "Giải pháp in tem nhãn tủ điện, viễn thông, đường ống và nhãn dán công nghiệp Brother.",
                "link": "https://www.brother.com.vn/",
                "isLocal": False
            }
        ]
    }

# --- 5. ZEBRA ---
def parse_zebra():
    ws = wb['SP ZEBRA']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'Máy in công nghiệp'
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        c2 = clean(r[2])
        c3 = clean(r[3])
        if is_stop_row(c0) or is_stop_row(c1) or is_stop_row(c2): break
        
        cat_hint = c1 or c0
        if cat_hint and ('máy in' in cat_hint.lower() or 'máy quét' in cat_hint.lower()):
            current_cat = cat_hint

        model_name = c3 or c2 or c1
        if not model_name or 'SERIAL' in model_name.upper() or 'MODEL' in model_name.upper(): continue

        status_raw = clean(r[4])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'hiếm' in status_raw.lower(): status = 'Hiếm'

        specs = clean(r[5])
        app = clean(r[6])
        suppliers = clean(r[7])
        web = clean(r[8])

        models = [m.strip() for m in re.split(r'[/,\n]+', model_name) if m.strip()]

        points = []
        if status_raw: points.append(f"**Đặc điểm**: {status_raw}")
        if specs: points.append(f"**Độ phân giải / Khổ in**: {specs}")
        if app: points.append(f"**Ứng dụng**: {app}")
        points.append(f"**Nhà cung cấp**: {suppliers or 'Việt Á, Baso, Tân Phát'}")

        products.append({
            "id": pid,
            "cat": current_cat,
            "subcat": specs or app or current_cat,
            "serial": f"Zebra {model_name.replace(chr(10), ' ')}",
            "status": status,
            "fake": 'Không',
            "renew": 'Không',
            "models": models,
            "replacement": "Dòng ZT231 thay thế ZT230; ZT411 thay thế ZT410" if 'ngừng' in status_raw.lower() else "",
            "points": points,
            "software": "ZebraDesigner 3",
            "brochure": web if web.startswith('http') else "https://www.zebra.com/",
            "suppliers": suppliers or "Việt Á, Baso, Tân Phát, Tân Hưng Hà"
        })
        pid += 1
    return {
        "brand": "Zebra",
        "products": products,
        "history": [
            {
                "category": "Máy In Mã Vạch Công Nghiệp Zebra",
                "steps": [
                    {
                        "era": "2000 - 2014",
                        "name": "S4M / ZM400 / 105SL",
                        "status": "Đã khai tử",
                        "badge": "discontinued",
                        "software": "ZebraDesigner v2",
                        "highlight": "Bộ khung kim loại bền bỉ nổi tiếng, hoạt động liên tục 24/7 trong các kho vận Logistics."
                    },
                    {
                        "era": "2014 - 2021",
                        "name": "ZT410 / ZT230 / GT800",
                        "status": "Thông dụng",
                        "badge": "classic",
                        "software": "ZebraDesigner 3",
                        "highlight": "Dòng máy in công nghiệp bán chạy nhất lịch sử Zebra, tốc độ in 14 ips, chuẩn mực ngành barcode."
                    },
                    {
                        "era": "2021 - Nay",
                        "name": "ZT411 / ZT231 / ZD421",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "Print DNA / Link-OS",
                        "highlight": "Màn hình cảm ứng màu 4.3 inch trực quan, hỗ trợ mã hóa RFID và kết nối đám mây từ xa."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "ZebraDesigner 3",
                "version": "v3.2+",
                "target": "Máy in mã vạch công nghiệp ZT series, máy in để bàn ZD series",
                "purpose": "Thiết kế tem nhãn mã vạch 1D, 2D QR Code, Datamatrix, liên kết cơ sở dữ liệu in tem hàng loạt.",
                "note": "Bản Standard miễn phí kèm theo máy, bản Professional hỗ trợ kết nối SQL/ODBC chuyên sâu.",
                "link": "https://www.zebra.com/us/en/support-downloads/printer-software/zebradesigner.html"
            }
        ],
        "brochures": [
            {
                "title": "Zebra Industrial Printers Portfolio Guide",
                "category": "Catalog Tổng Hợp",
                "desc": "Tổng hợp các dòng máy in tem nhãn mã vạch công nghiệp ZT111, ZT231, ZT411, ZT610.",
                "link": "https://www.zebra.com/us/en/products/printers/industrial.html",
                "isLocal": False
            }
        ]
    }

# --- 6. PRO-FACE & HMI ---
def parse_hmi():
    ws = wb['SP HMI']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        if is_stop_row(c0) or is_stop_row(c1): break
        if not c1 or 'MODEL' in c1.upper(): continue

        brand_hint = c0 or 'PROFACE'
        model_name = c1
        size = clean(r[2])
        status_raw = clean(r[3])
        status = 'Thông dụng'
        if 'ngừng' in status_raw.lower(): status = 'Ngừng sx'
        elif 'hiếm' in status_raw.lower(): status = 'Hiếm'

        fake = 'Có' if 'yes' in clean(r[4]).lower() else 'Không'
        renew = 'Có' if 'yes' in clean(r[5]).lower() else 'Không'
        suppliers = clean(r[6])
        web = clean(r[7])
        notes = clean(r[8]) if len(r) > 8 else ''

        points = []
        if size: points.append(f"**Kích thước màn hình**: {size}")
        if status_raw: points.append(f"**Đặc điểm**: {status_raw}")
        if fake == 'Có': points.append("⚠️ Cảnh báo thực chiến: Thị trường có hàng FAKE bo mạch nhái.")
        if renew == 'Có': points.append("🔄 Cảnh báo thực chiến: Có rủi ro hàng Renew / ép lại cảm ứng tháo máy.")
        if notes: points.append(f"💡 {notes}")
        points.append(f"**Nhà cung cấp**: {suppliers or 'Hợp Long, Nhập khẩu'}")

        products.append({
            "id": pid,
            "cat": f"Màn hình {brand_hint}",
            "subcat": f"Màn hình cảm ứng HMI {size}" if size else f"Màn hình HMI {brand_hint}",
            "serial": f"{brand_hint} {model_name}",
            "status": status,
            "fake": fake,
            "renew": renew,
            "models": [model_name],
            "replacement": "Chuyển sang dòng GP4000 hoặc ST6000 thế hệ mới" if status == 'Ngừng sx' else "",
            "points": points,
            "software": "GP-Pro EX",
            "brochure": web if web.startswith('http') else "https://www.proface.com/",
            "suppliers": suppliers or "Hợp Long, Nhập khẩu"
        })
        pid += 1
    return {
        "brand": "Pro-face",
        "products": products,
        "history": [
            {
                "category": "Màn Hình Cảm Ứng HMI Pro-face",
                "steps": [
                    {
                        "era": "1998 - 2012",
                        "name": "GP2000 / GP3000 Series",
                        "status": "Đã khai tử",
                        "badge": "discontinued",
                        "software": "GP-PRO/PB III",
                        "highlight": "Thương hiệu màn hình cảm ứng HMI tiên phong số 1 thế giới từ Digital Electronics Nhật Bản."
                    },
                    {
                        "era": "2013 - Nay",
                        "name": "GP4000 Series (GP4301, GP4401, GP4501)",
                        "status": "Thông dụng",
                        "badge": "classic",
                        "software": "GP-Pro EX v4+",
                        "highlight": "Dòng HMI tiêu chuẩn hàng đầu trong các máy móc Nhật Bản, châu Âu."
                    },
                    {
                        "era": "2020 - Nay",
                        "name": "ST6000 / ET6000 Series",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "BLUE / GP-Pro EX",
                        "highlight": "Viền mỏng hiện đại, độ phân giải cao True Color, hỗ trợ giao diện Web HMI HTML5."
                    }
                ]
            }
        ],
        "software": [
            {
                "name": "GP-Pro EX",
                "version": "v4.09+",
                "target": "Màn hình HMI Pro-face GP4000, GP3000, ST6000, SP5000",
                "purpose": "Thiết kế giao diện HMI, cấu hình truyền thông đa giao thức kết nối đồng thời nhiều hãng PLC.",
                "note": "Phần mềm HMI mạnh mẽ và linh hoạt nhất ngành tự động hóa, hỗ trợ mô phỏng Offline Simulation.",
                "link": "https://www.proface.com/en/product/software/gpproex"
            }
        ],
        "brochures": [
            {
                "title": "Pro-face HMI & Industrial Computer General Catalog",
                "category": "Catalog Tổng Hợp",
                "desc": "Tổng hợp dải sản phẩm màn hình cảm ứng HMI, máy tính công nghiệp IPC Pro-face Schneider Electric.",
                "link": "https://www.proface.com/",
                "isLocal": False
            }
        ]
    }

# --- 7. AN TOÀN XE NÂNG & CẦU TRỤC ---
def parse_xenang():
    ws = wb['SP Xe Nâng & Cầu Trục']
    rows = list(ws.iter_rows(values_only=True))
    products = []
    current_cat = 'Đèn vạch vùng an toàn'
    pid = 1
    for r in rows[1:]:
        c0 = clean(r[0])
        c1 = clean(r[1])
        if is_stop_row(c0) or is_stop_row(c1): break
        if not c1:
            if c0: current_cat = c0
            continue
        if c0: current_cat = c0

        model_name = c1
        status_raw = clean(r[2])
        status = 'Thông dụng'
        if 'không thông dụng' in status_raw.lower() or 'hiếm' in status_raw.lower(): status = 'Hiếm'
        elif 'ngừng' in status_raw.lower(): status = 'Ngừng sx'

        models_raw = clean(r[5])
        models = [m.strip() for m in re.split(r'[,;\n]+', models_raw) if m.strip()]
        if not models: models = [model_name]

        suppliers = clean(r[6])

        points = []
        if models: points.append(f"**Mã sản phẩm**: {', '.join(models)}")
        points.append(f"**Chuyên dụng**: Lắp đặt an toàn cho xe nâng hàng, cầu trục công xưởng")
        points.append(f"**Nhà cung cấp**: {suppliers or 'DACO, VTH, VIETMRO'}")

        products.append({
            "id": pid,
            "cat": current_cat.replace('\n', ' '),
            "subcat": "Hệ thống cảnh báo an toàn nhà xưởng DACO",
            "serial": model_name.replace('\n', ' '),
            "status": status,
            "fake": 'Không',
            "renew": 'Không',
            "models": models,
            "replacement": "",
            "points": points,
            "software": "Không dùng",
            "brochure": "https://daco.vn/",
            "suppliers": suppliers or "DACO, VTH, VIETMRO"
        })
        pid += 1
    return {
        "brand": "An Toàn Xe Nâng & Cầu Trục",
        "products": products,
        "history": [
            {
                "category": "Hệ Thống An Toàn Xe Nâng & Cầu Trục",
                "steps": [
                    {
                        "era": "2010 - 2018",
                        "name": "Còi Lùi & Đèn Chớp Cơ Bản",
                        "status": "Kinh điển",
                        "badge": "classic",
                        "software": "Đấu nối nguồn 12-48V",
                        "highlight": "Các giải pháp cảnh báo truyền thống dùng còi lùi tít tít và đèn chớp gắn nóc xe nâng."
                    },
                    {
                        "era": "2018 - Nay",
                        "name": "Đèn Vạch Vùng & Đèn Rọi Điểm LED",
                        "status": "Chuẩn hiện hành",
                        "badge": "current",
                        "software": "Quang học LED",
                        "highlight": "Chiếu vạch LED đỏ/xanh bao quanh thân xe nâng và điểm rọi cảnh báo trước mũi xe 3-5m."
                    },
                    {
                        "era": "2022 - Nay",
                        "name": "Camera AI Nhận Diện Người 360",
                        "status": "Thế hệ mới",
                        "badge": "future",
                        "software": "AI Computer Vision",
                        "highlight": "Camera AI phân biệt chính xác con người, tự động phát âm thanh cảnh báo và kích hoạt phanh."
                    }
                ]
            }
        ],
        "software": [],
        "brochures": [
            {
                "title": "DACO Forklift Safety Solutions Catalog",
                "category": "Catalog Sản Phẩm",
                "desc": "Tổng hợp giải pháp đèn vạch an toàn, camera AI phát hiện người đi bộ và hệ thống chống va chạm xe nâng.",
                "link": "https://daco.vn/",
                "isLocal": False
            }
        ]
    }

# Execute parsers
brands_data = {
    'omron': parse_omron(),
    'autonics': parse_autonics(),
    'patlite': parse_patlite(),
    'brother': parse_brother(),
    'zebra': parse_zebra(),
    'proface': parse_hmi(),
    'xenang': parse_xenang()
}

out_dir = 'e:/sp/Mitutoyo/cam_nang_san_pham/data'

# Meta information for BRANDS list
new_brand_metas = [
    {
        "id": "omron",
        "name": "Omron",
        "code": "OMRON",
        "badge": "Tự Động Hóa & Cảm Biến",
        "color": "#005bac",
        "icon": "🏢",
        "tagline": "PLC, Biến Tần, Cảm Biến Tiệm Cận & Quang, Rơ Le Trung Gian, Đồng Hồ Nhiệt Độ E5CC",
        "productCount": len(brands_data['omron']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "autonics",
        "name": "Autonics",
        "code": "AUTONICS",
        "badge": "Cảm Biến & Điều Khiển",
        "color": "#0284c7",
        "icon": "🏢",
        "tagline": "Cảm Biến Quang & Tiệm Cận, Đồng Hồ Nhiệt Độ TK/TC, Bộ Đếm Timer, Encoder Vòng Quay",
        "productCount": len(brands_data['autonics']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "patlite",
        "name": "Patlite",
        "code": "PATLITE",
        "badge": "Đèn Tháp & Còi Báo Động",
        "color": "#dc2626",
        "icon": "🏢",
        "tagline": "Đèn Tháp Tín Hiệu LR Series, Đèn Báo Quay & Nhấp Nháy, Còi Báo Động Nhà Xưởng EHS",
        "productCount": len(brands_data['patlite']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "brother",
        "name": "Brother",
        "code": "BROTHER",
        "badge": "Máy In Nhãn & In Ống",
        "color": "#2563eb",
        "icon": "🏢",
        "tagline": "Máy In Nhãn Cầm Tay P-Touch, Máy In Ống Đầu Cốt PT-E850TKW, Máy In Nhãn Công Nghiệp",
        "productCount": len(brands_data['brother']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "zebra",
        "name": "Zebra",
        "code": "ZEBRA",
        "badge": "Máy In Mã Vạch & Barcode",
        "color": "#334155",
        "icon": "🏢",
        "tagline": "Máy In Mã Vạch Công Nghiệp ZT Series, Máy In Để Bàn ZD Series, Máy Quét Barcode 1D/2D",
        "productCount": len(brands_data['zebra']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "proface",
        "name": "Pro-face",
        "code": "PROFACE",
        "badge": "Màn Hình Cảm Ứng HMI",
        "color": "#059669",
        "icon": "🏢",
        "tagline": "Màn Hình HMI Cảm Ứng GP4000, ST6000, ET6000, Màn Hình Weintek & Beijer",
        "productCount": len(brands_data['proface']['products']),
        "hasSoftware": True,
        "hasHistory": True
    },
    {
        "id": "xenang",
        "name": "An Toàn Xe Nâng & Cầu Trục",
        "code": "SAFETY",
        "badge": "An Toàn Nhà Xưởng",
        "color": "#d97706",
        "icon": "🏢",
        "tagline": "Đèn Vạch Vùng An Toàn LED, Camera AI Nhận Diện Người, Cảnh Báo Tốc Độ Xe Nâng",
        "productCount": len(brands_data['xenang']['products']),
        "hasSoftware": False,
        "hasHistory": True
    }
]

# Write JS and JSON files for each brand
for bid, bdata in brands_data.items():
    var_name = f"PORTAL_DATA_{bid.upper()}"
    js_content = f"window.{var_name} = " + json.dumps(bdata, ensure_ascii=False, indent=2) + ";\n"
    with open(os.path.join(out_dir, f"{bid}.js"), 'w', encoding='utf-8') as f:
        f.write(js_content)
    with open(os.path.join(out_dir, f"{bid}.json"), 'w', encoding='utf-8') as f:
        json.dump(bdata, f, ensure_ascii=False, indent=2)
    print(f"Generated {bid}.js and {bid}.json ({len(bdata['products'])} products)")

# Load existing brands.json
with open(os.path.join(out_dir, 'brands.json'), 'r', encoding='utf-8') as f:
    existing_brands = json.load(f)

# Keep mitsubishi, mitutoyo, qlight, and append new brands if not exists
existing_ids = {b['id'] for b in existing_brands}
all_brands = [b for b in existing_brands if b['id'] in ['mitsubishi', 'mitutoyo', 'qlight']]
for nb in new_brand_metas:
    all_brands.append(nb)

# Write updated brands.json and brands.js
with open(os.path.join(out_dir, 'brands.json'), 'w', encoding='utf-8') as f:
    json.dump(all_brands, f, ensure_ascii=False, indent=2)

brands_js_content = "window.PORTAL_BRANDS = " + json.dumps(all_brands, ensure_ascii=False, indent=2) + ";\n"
with open(os.path.join(out_dir, 'brands.js'), 'w', encoding='utf-8') as f:
    f.write(brands_js_content)

print(f"Updated brands.js and brands.json with total {len(all_brands)} brands.")
